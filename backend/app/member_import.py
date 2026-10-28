import csv
import io
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import ProjectMember, Role, User
from .schemas import BulkImportRowInput

HEADERS = ['工号', '用户名', '所属部门', '项目角色', '邮箱']
FIELD_NAMES = ['employee_id', 'username', 'department', 'role_name', 'email']
RESERVED_ROLES = {'系统管理员', '初始化用户'}
MAX_IMPORT_ROWS = 5000


def _cell(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_import_file(content: bytes, filename: str) -> list[dict]:
    suffix = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    try:
        if suffix == 'xlsx':
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            raw_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        elif suffix == 'csv':
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = content.decode('gb18030')
            raw_rows = list(csv.reader(io.StringIO(text)))
        else:
            raise HTTPException(422, '仅支持 .xlsx 或 .csv 文件')
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, '文件无法解析，请使用系统下载的 Excel 模板') from None
    raw_rows = [row for row in raw_rows if any(_cell(value) for value in row)]
    if not raw_rows:
        raise HTTPException(422, '导入文件为空')
    headers = [_cell(value) for value in raw_rows[0]]
    if len(headers) != len(set(headers)) or set(headers) - set(HEADERS) or not {'工号', '用户名'}.issubset(headers):
        raise HTTPException(422, f'表头不正确，应包含：{"、".join(HEADERS)}')
    data_rows = raw_rows[1:]
    if not data_rows:
        raise HTTPException(422, '模板中没有用户数据')
    if len(data_rows) > MAX_IMPORT_ROWS:
        raise HTTPException(422, f'单次最多导入 {MAX_IMPORT_ROWS} 名用户')
    result = []
    for values in data_rows:
        row = {FIELD_NAMES[HEADERS.index(header)]: _cell(values[index]) if index < len(values) else ''
               for index, header in enumerate(headers)}
        row['row_number'] = len(result) + 2
        result.append(row)
    return result


def validate_import_rows(db: Session, project_id: str, default_role_id: str, raw_rows: list[dict]) -> dict:
    roles = list(db.scalars(select(Role)).all())
    roles_by_id = {role.id: role for role in roles}
    roles_by_name = {role.name: role for role in roles}
    default_role = roles_by_id.get(default_role_id)
    if not default_role or default_role.name in RESERVED_ROLES:
        raise HTTPException(422, '默认角色不存在或不允许分配')

    normalized = []
    seen_employee_ids = set()
    for raw in raw_rows:
        errors = []
        try:
            row = BulkImportRowInput.model_validate(raw)
            values = row.model_dump()
        except ValidationError as exc:
            values = dict(raw)
            errors = ['；'.join(error['msg'] for error in exc.errors())]
            row = None
        employee_id = _cell(values.get('employee_id'))
        if employee_id:
            if employee_id in seen_employee_ids:
                errors.append('工号在文件中重复')
            seen_employee_ids.add(employee_id)
        normalized.append({'raw': values, 'row': row, 'errors': errors})

    employee_ids = [item['row'].employee_id for item in normalized if item['row']]
    users = {user.employee_id: user for user in db.scalars(select(User).where(User.employee_id.in_(employee_ids))).all()} if employee_ids else {}
    memberships = {}
    if users:
        rows = db.scalars(select(ProjectMember).where(ProjectMember.project_id == project_id,
                                                      ProjectMember.user_id.in_([user.id for user in users.values()]))).all()
        memberships = {member.user_id: member for member in rows}

    output_rows = []
    for item in normalized:
        raw, row, errors = item['raw'], item['row'], list(item['errors'])
        user = users.get(row.employee_id) if row else None
        role = default_role
        if row and row.role_name:
            role = roles_by_name.get(row.role_name)
            if not role:
                errors.append('项目角色不存在')
        if role and role.name in RESERVED_ROLES:
            errors.append(f'{role.name}不允许通过批量导入分配')
        if user and not user.is_active:
            errors.append('平台用户已停用')
        if user and user.is_superuser:
            errors.append('系统管理员不能通过批量导入加入项目')
        membership = memberships.get(user.id) if user else None
        action = 'EXISTING_MEMBER' if membership else 'ADD_MEMBER' if user else 'CREATE_USER_AND_MEMBER'
        if membership:
            role = roles_by_id.get(membership.role_id, role)
        output_rows.append({
            'row_number': raw.get('row_number'), 'employee_id': raw.get('employee_id', ''),
            'username': raw.get('username', ''), 'department': raw.get('department') or None,
            'email': raw.get('email') or None, 'role_id': role.id if role else None,
            'role_name': role.name if role else raw.get('role_name'), 'action': action,
            'status': 'ERROR' if errors else 'READY', 'errors': errors,
        })
    error_count = sum(row['status'] == 'ERROR' for row in output_rows)
    summary = {
        'total': len(output_rows), 'valid': len(output_rows) - error_count, 'errors': error_count,
        'new_users': sum(row['status'] == 'READY' and row['action'] == 'CREATE_USER_AND_MEMBER' for row in output_rows),
        'new_members': sum(row['status'] == 'READY' and row['action'] in {'CREATE_USER_AND_MEMBER', 'ADD_MEMBER'} for row in output_rows),
        'existing_members': sum(row['status'] == 'READY' and row['action'] == 'EXISTING_MEMBER' for row in output_rows),
    }
    return {'summary': summary, 'rows': output_rows, '_users': users, '_memberships': memberships}


def apply_import(db: Session, project_id: str, default_role_id: str, raw_rows: list[dict]) -> dict:
    validated = validate_import_rows(db, project_id, default_role_id, raw_rows)
    if validated['summary']['errors']:
        details = [f"第 {row['row_number']} 行：{'；'.join(row['errors'])}" for row in validated['rows'] if row['errors']][:5]
        raise HTTPException(422, f"导入校验失败：{'；'.join(details)}")
    users = validated['_users']
    for row in validated['rows']:
        if row['action'] == 'CREATE_USER_AND_MEMBER':
            user = User(employee_id=row['employee_id'], username=row['username'], department=row['department'],
                        email=row['email'], is_active=True, is_superuser=False)
            db.add(user)
            users[row['employee_id']] = user
    db.flush()
    for row in validated['rows']:
        if row['action'] != 'EXISTING_MEMBER':
            db.add(ProjectMember(project_id=project_id, user_id=users[row['employee_id']].id, role_id=row['role_id']))
    return {'summary': validated['summary'], 'rows': validated['rows']}


def import_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '用户初始化模板'
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='315EFB')
        cell.alignment = Alignment(horizontal='center')
    for column, width in zip('ABCDE', [20, 28, 28, 22, 36]):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    guide = workbook.create_sheet('填写说明')
    instructions = [
        ['字段', '是否必填', '说明'],
        ['工号', '必填', '全局唯一，作为用户匹配和初始化的唯一业务标识'],
        ['用户名', '必填', '用户展示名称，允许重复'],
        ['所属部门', '选填', '仅用于用户档案展示和筛选'],
        ['项目角色', '选填', '为空时使用导入页面选择的默认角色；已有成员保留原角色'],
        ['邮箱', '选填', '普通联系信息'],
        ['注意', '', '不要修改表头；单次最多 5000 行；系统管理员和初始化用户角色不可导入'],
    ]
    for row in instructions:
        guide.append(row)
    guide.column_dimensions['A'].width = 18
    guide.column_dimensions['B'].width = 14
    guide.column_dimensions['C'].width = 90
    for cell in guide[1]:
        cell.font = Font(bold=True)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
